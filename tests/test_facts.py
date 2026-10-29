"""Tables and typed facts: what the readers find in rows and columns, and what the
fact rules make of each cell. Inputs are written here, so each expected fact is visible."""
from pathlib import Path

from openpyxl import Workbook

from engineering_search.core import Table, TableRow
from engineering_search.indexing import Indexer
from engineering_search.indexing.facts import extract_facts
from engineering_search.indexing.readers import CsvReader, MarkdownReader, XlsxReader
from engineering_search.indexing.status import INDEXED_PARTIAL
from engineering_search.store import SQLiteSearchIndex

SCHEDULE = """\
Component,Material grade,Allowable UDL (kN/m²),Design imposed load (kN/m²),Status
B-104,S355J2,4.5,,Approved
F-021,C30/37,,3.2,Approved
"""


def facts_of(table: Table) -> dict[tuple[str, str], tuple]:
    return {(f.subject, f.parameter): (f.value, f.role, f.si_value, f.dimension) for f in extract_facts((table,))}


def test_a_csv_row_gives_a_table_row_and_a_passage_that_cites_its_line():
    extraction = CsvReader().read("s.csv", "v", SCHEDULE.encode("utf-8-sig"))
    table = extraction.tables[0]
    assert table.columns[2] == "Allowable UDL (kN/m²)"
    assert table.rows[0] == TableRow(1, ("B-104", "S355J2", "4.5", "", "Approved"), 2)
    first = extraction.passages[0]
    assert (first.start_line, first.text) == (2, "B-104,S355J2,4.5,,Approved")
    assert first.heading.startswith("Component,Material grade")  # Column names are searchable.


def test_capacity_and_demand_are_different_parameters():
    table = CsvReader().read("s.csv", "v", SCHEDULE.encode()).tables[0]
    facts = facts_of(table)
    assert facts[("B-104", "material_grade")][:2] == ("S355J2", "specification")
    assert facts[("B-104", "allowable_distributed_load")] == ("4.5", "capacity", 4500.0, "pressure")
    assert facts[("F-021", "imposed_load")] == ("3.2", "demand", 3200.0, "pressure")
    assert ("F-021", "allowable_distributed_load") not in facts
    assert facts[("B-104", "compliance_status")][:2] == ("Approved", "status")


def test_units_are_normalised_but_never_rescaled_by_intent():
    rows = (TableRow(1, ("P-1", "5 N/m²"), 2), TableRow(2, ("P-2", "5 kN/m²"), 3), TableRow(3, ("P-3", "-12 mm"), 4))
    facts = extract_facts((Table(1, "Probe", ("Member", "Allowable load"), rows),))
    assert [(f.subject, f.si_value, f.dimension) for f in facts] == [
        ("P-1", 5.0, "pressure"), ("P-2", 5000.0, "pressure"), ("P-3", -0.012, "length")]


def test_a_key_value_table_takes_the_parameter_from_the_row_label():
    markdown = ("# Synthetic project\n\n## Roof load\n\n| Parameter | Value | Approval status |\n"
                "| --- | --- | --- |\n| Maximum roof load | 7 kN/m² | Approved |\n")
    extraction = MarkdownReader().read("p.md", "v", markdown.encode())
    assert extraction.tables[0].rows[0].line == 7
    facts = facts_of(extraction.tables[0])
    assert facts[("Roof load", "allowable_distributed_load")] == ("7 kN/m²", "capacity", 7000.0, "pressure")
    assert facts[("Roof load", "compliance_status")][0] == "Approved"


def test_a_pipe_table_inside_a_code_fence_is_not_a_table():
    markdown = "# A\n\n```\n| a | b |\n|---|---|\n| 1 | 2 |\n```\n"
    assert MarkdownReader().read("a.md", "v", markdown.encode()).tables == ()


def test_a_project_column_overrides_the_document_for_each_row():
    rows = (TableRow(1, ("Cedar", "CB-12", "S355"), 2), TableRow(2, ("Harbour", "B-104", "S355J2"), 3))
    facts = extract_facts((Table(1, "Register", ("Project", "Component", "Material grade"), rows),))
    assert [(f.project, f.subject, f.value) for f in facts] == [("Cedar", "CB-12", "S355"),
                                                                ("Harbour", "B-104", "S355J2")]


def workbook(path: Path, rows: list[list]) -> Path:
    book = Workbook()
    book.active.title = "Roof members"
    for row in rows:
        book.active.append(row)
    book.save(path)
    return path


def test_a_formula_without_a_cached_value_is_left_empty_and_reported(tmp_path):
    path = workbook(tmp_path / "roof.xlsx", [["Member", "Material grade", "Allowable load (kN/m²)"],
                                             ["R-310", "S355", 3.8], ["R-313", "S355", "=3.5*1"]])
    extraction = XlsxReader().read("roof.xlsx", "v", path.read_bytes())
    assert extraction.status == INDEXED_PARTIAL
    assert extraction.warnings == ("formulas without a cached value, left empty: Roof members!C3",)
    assert [r.cells for r in extraction.tables[0].rows] == [("R-310", "S355", "3.8"), ("R-313", "S355", "")]
    first = extraction.passages[0]
    assert (first.page, first.start_line, first.text) == (1, 2, "R-310 | S355 | 3.8")
    facts = facts_of(extraction.tables[0])
    assert ("R-313", "allowable_distributed_load") not in facts  # Not evaluated, so no value.
    assert facts[("R-310", "allowable_distributed_load")][2] == 3800.0


def test_tables_and_facts_are_stored_with_the_document_and_removed_with_it(tmp_path):
    (tmp_path / "s.csv").write_text(SCHEDULE)
    indexer, index = Indexer(tmp_path), SQLiteSearchIndex()
    plan = indexer.plan({})
    index.apply(plan.upserts, plan.removals)
    stored = index.structured({"s.csv"})
    assert [r[4] for r in stored["facts"] if r[5] == "allowable_distributed_load"] == ["B-104"]
    assert len(stored["rows"]) == 2 and stored["tables"][0][2] == "s.csv"
    assert index.structured({"other.csv"})["facts"] == []  # Only the files asked for.
    index.remove("s.csv")
    assert index.structured({"s.csv"}) == {"documents": [], "tables": [], "rows": [], "facts": []}
